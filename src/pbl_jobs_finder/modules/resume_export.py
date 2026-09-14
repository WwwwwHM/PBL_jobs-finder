"""Export an editable, restrained Word resume from optimized Markdown text."""

from __future__ import annotations

import re
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

_INVALID_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)、]\s+(.*)$")


class ResumeExportError(RuntimeError):
    """The optimized resume could not be written as a Word document."""


def create_resume_docx(
    optimized_text: str,
    target_position: str,
    output_dir: str | Path,
    *,
    record_id: int | None = None,
) -> Path:
    """Create a polished DOCX while preserving the user's editable text."""

    text = (optimized_text or "").strip()
    if not text:
        raise ValueError("优化后的简历内容不能为空")

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    stem = _safe_filename(target_position) or "新版简历"
    suffix = str(record_id) if record_id is not None else uuid4().hex[:8]
    destination = output_path / f"{stem}_优化简历_{suffix}.docx"

    document = Document()
    _configure_document(document)
    _append_markdown(document, text)
    document.core_properties.title = f"{target_position} 优化简历"
    document.core_properties.subject = target_position
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    try:
        document.save(destination)
    except (OSError, ValueError) as exc:
        raise ResumeExportError("Word 简历生成失败，请稍后重试") from exc
    return destination


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    title = document.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _remove_paragraph_borders(title)

    for style_name, size in (("Heading 1", 13), ("Heading 2", 11)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True


def _append_markdown(document: Document, text: str) -> None:
    if not any(line.strip().startswith("# ") for line in text.splitlines()):
        document.add_paragraph("个人简历", style="Title")
    has_title = False
    previous_was_title = False
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("### "):
            document.add_paragraph(_clean_inline(line[4:]), style="Heading 2")
            previous_was_title = False
            continue
        if line.startswith("## "):
            document.add_paragraph(_clean_inline(line[3:]), style="Heading 1")
            previous_was_title = False
            continue
        if line.startswith("# "):
            style = "Title" if not has_title else "Heading 1"
            document.add_paragraph(_clean_inline(line[2:]), style=style)
            has_title = True
            previous_was_title = style == "Title"
            continue
        bullet = _BULLET.match(line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _append_inline_runs(paragraph, bullet.group(1))
            previous_was_title = False
            continue
        numbered = _NUMBERED.match(line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            _append_inline_runs(paragraph, numbered.group(1))
            previous_was_title = False
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        if previous_was_title:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
        _append_inline_runs(paragraph, line)
        previous_was_title = False


def _append_inline_runs(paragraph: Paragraph, line: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", line)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)


def _clean_inline(value: str) -> str:
    return value.replace("**", "").strip()


def _safe_filename(value: str) -> str:
    cleaned = _INVALID_FILENAME.sub("_", (value or "").strip())
    return cleaned.rstrip(" .")[:50]


def _remove_paragraph_borders(style: object) -> None:
    paragraph_properties = style.element.get_or_add_pPr()
    existing = paragraph_properties.find(qn("w:pBdr"))
    if existing is not None:
        paragraph_properties.remove(existing)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    borders.append(bottom)
    paragraph_properties.append(borders)


__all__ = ["ResumeExportError", "create_resume_docx"]
